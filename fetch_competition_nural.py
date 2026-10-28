#!/usr/bin/env python3
"""
fetch_competition_nural.py
--------------------------
Pulls the Nural **Competition Sales Report** for today (FTD) and sends a
TSI-wise Telegram message listing the East Reliance ISDs who have **NOT**
reported any competition sale today.

Runs 6 PM & 9 PM IST via GitHub Actions (competition_report.yml).

Pipeline:
  1. Playwright logs into motoindia.nuralsales.app (same portal / login as the
     attendance scraper) and exports the Competition Sales Report for today only
     (FROM = TO = today) as Excel — reuses the proven date-pick + export + queue
     flow from fetch_weekoff_nural.py.
  2. Parse the export. The on-screen report identifies the ISD by NAME
     (ISPNAME), so we join to the East Reliance roster by ISP Code when the
     export carries one, else by normalised ISP Name.
  3. "Reported" = the ISD has >= 1 competition-sale row today. "Not reported" =
     every ACTIVE roster ISD not in that set, excluding anyone on Weekly-Off /
     Leave today (from the Attendance Tracker WeekOff Log) and Inactive ISDs.
  4. Group the not-reported ISDs by TSI and send one Telegram message to the
     same chat as the other reports (TELEGRAM_CHAT_ID).

Env:
  PORTAL_USER / PORTAL_PASS   Nural login (reused from the sales repo).
  TARGETS_CSV_URL             published Targets CSV (roster / TSI join).
  WEEKOFF_CSV_URL             published Attendance Tracker WeekOff Log CSV
                              (to exclude Weekly-Off / Leave ISDs for today).
  TELEGRAM_BOT_TOKEN / TELEGRAM_CHAT_ID
Optional:
  COMP_XLSX        local .xlsx path -> skip Playwright, parse this (offline).
  TODAY_OVERRIDE   YYYY-MM-DD -> anchor on this date instead of the wall clock.
  SEND_TELEGRAM=0  build the message but print instead of sending.
  HEADLESS=false   show the browser (local debugging).
"""

import os
import re
import csv
import html
import urllib.request
from datetime import datetime, timedelta, date, timezone
from pathlib import Path

import fetch_weekoff_nural as WO        # proven Nural login / date-pick / queue / telegram
import attendance_reports as AR         # parse_log / excluded_codes / _is_inactive

IST = timezone(timedelta(hours=5, minutes=30))
COMP_URL = WO.NURAL_ROOT + "/competition-sales-report"

# Attendance Tracker · WeekOff Log tab (date-stamped attendance log) — same
# published CSV the Remarks pipeline reads. Env override wins.
WEEKOFF_CSV_URL = os.environ.get("WEEKOFF_CSV_URL") or \
    "https://docs.google.com/spreadsheets/d/1BsShTlrnwYGBUKMcURMI1Qla-pnttWFrPbMmiLNu6Eo/export?format=csv&gid=1706495387"


def env(n, d=""):
    return (os.environ.get(n) or d).strip()

def esc(s):
    return html.escape(str(s or ""), quote=True)

def _norm_name(s):
    """Lower-case, strip punctuation, collapse spaces — for name matching."""
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", str(s or "").lower())).strip()

# ── Nural export (Competition Sales Report) ──────────────────────────────────

def _trigger_export(page):
    """Click whatever control downloads the result grid. The competition report
    uses an 'Export' button (not the attendance exceldownload icon), so try a
    few selectors and let the caller handle direct / download / queue."""
    candidates = [
        'button:has-text("Export")',
        'img[src*="exceldownload"]',
        'img[alt*="download" i]',
    ]
    for sel in candidates:
        loc = page.locator(sel).first
        try:
            if loc.count() and loc.is_visible():
                loc.scroll_into_view_if_needed()
                loc.evaluate("el => el.click()")
                return True
        except Exception:
            continue
    return False

def _export_competition(page, from_dt, to_dt, out_path):
    page.goto(COMP_URL, wait_until="networkidle", timeout=60000)
    if "/login" in page.url.lower():
        raise RuntimeError("Redirected to /login opening the competition report")
    page.wait_for_selector(WO.CAL_ICON, state="visible", timeout=45000)
    WO._set_date(page, 0, from_dt)
    WO._set_date(page, 1, to_dt)

    direct, downloads = [], []
    page.on("response", lambda r: r.url.lower().endswith(".xlsx") and direct.append(r))
    page.on("download", lambda d: downloads.append(d))
    started = datetime.now(IST)

    # SEARCH carries an icon, so match by visible text, not accessible name.
    btn = page.locator("button").filter(has_text=re.compile(r"^\s*SEARCH\s*$")).first
    btn.scroll_into_view_if_needed()
    btn.click()

    # Wait for the grid to render (TOTAL RECORDS footer), then trigger export.
    try:
        page.get_by_text(re.compile(r"TOTAL RECORDS", re.I)).first.wait_for(timeout=60000)
    except Exception:
        pass
    _trigger_export(page)

    deadline = datetime.now(IST) + timedelta(seconds=120)
    while datetime.now(IST) < deadline:
        if downloads:
            downloads[0].save_as(out_path); print("[COMP] downloaded", out_path); return
        if direct:
            body = direct[0].body()
            if len(body) >= 1024:
                Path(out_path).write_bytes(body); print("[COMP] downloaded(direct)", out_path); return
        if page.get_by_text(re.compile(
                r"status of your request|report queue|added to queue|submitted successfully",
                re.I)).count():
            break
        # Re-attempt the export trigger in case the first click predated the grid.
        _trigger_export(page)
        page.wait_for_timeout(1500)
    WO._download_from_queue(page, started, out_path, "competition")

def comp_download(day, out_dir):
    from playwright.sync_api import sync_playwright
    user = env("PORTAL_USER") or env("NURAL_USERNAME")
    pw   = env("PORTAL_PASS") or env("NURAL_PASSWORD")
    if not (user and pw):
        raise RuntimeError("Set PORTAL_USER / PORTAL_PASS (Nural login).")
    out_dir = Path(out_dir); out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "competition.xlsx"
    headless = env("HEADLESS", "true").lower() != "false"
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=headless, args=["--no-sandbox"])
        ctx = browser.new_context(accept_downloads=True, timezone_id="Asia/Kolkata")
        page = ctx.new_page()
        try:
            WO._login(page, user, pw)
            _export_competition(page, day, day, str(out))
        except Exception:
            try:
                (out_dir / "debug.html").write_text(page.content(), encoding="utf-8")
                page.screenshot(path=str(out_dir / "debug.png"), full_page=True)
            except Exception:
                pass
            raise
        finally:
            ctx.close(); browser.close()
    return out

# ── xlsx parsing (header-tolerant) ───────────────────────────────────────────

def _read_comp_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    def nrm(s):
        return re.sub(r"[\s_]+", "", str(s or "")).lower()
    anchors = {"ispname", "ispcode", "entitycode", "usercode", "saledate"}
    hdr_i = 0
    for i, r in enumerate(rows[:12]):
        if {nrm(c) for c in r} & anchors:
            hdr_i = i; break
    header = [str(c or "").strip() for c in rows[hdr_i]]
    out = []
    for r in rows[hdr_i + 1:]:
        if not any(str(c or "").strip() for c in r):
            continue
        out.append({header[j]: r[j] for j in range(min(len(header), len(r)))})
    return out

def reporters_for(day, rows, codes):
    """Set of roster ISP Codes that reported >=1 competition sale on `day`.
    Prefers an ISP-code column; else matches by normalised ISP Name."""
    if not rows:
        return set(), "none"
    r0 = rows[0]
    k_code = WO._key(r0, "ISP Code", "ISPCode", "EntityCode", "User Code")
    k_name = WO._key(r0, "ISPNAME", "ISP Name", "ISPName")
    k_date = WO._key(r0, "Sale Date", "SaleDate", "Date")

    name_to_codes = {}
    for c, info in codes.items():
        nm = _norm_name(info.get("isd"))
        if nm:
            name_to_codes.setdefault(nm, set()).add(c)

    reported = set()
    for r in rows:
        if k_date:
            d = WO.parse_any_date(r.get(k_date))
            if d and d != day:
                continue
        if k_code:
            c = str(r.get(k_code) or "").strip()
            if c in codes:
                reported.add(c)
        elif k_name:
            for c in name_to_codes.get(_norm_name(r.get(k_name)), ()):
                reported.add(c)
    return reported, ("code" if k_code else "name")

# ── Weekly-Off / Leave exclusion for today ───────────────────────────────────

def excluded_today(day):
    """ISP Codes on Weekly-Off / Leave on `day`, from the Attendance Tracker
    WeekOff Log. Best-effort — returns empty set on any failure."""
    if not WEEKOFF_CSV_URL:
        return set()
    try:
        rows = WO._fetch_csv(WEEKOFF_CSV_URL)
        log = AR.parse_log(rows)
        excl = AR.excluded_codes(log, day)
        print("[COMP] WeekOff/Leave excluded today:", len(excl))
        return excl
    except Exception as e:
        print("[COMP] WeekOff fetch failed (non-fatal):", str(e)[:120])
        return set()

# ── Telegram message ─────────────────────────────────────────────────────────

def build_message(not_reported, codes, day, now_ist, match_mode):
    by_tsi = {t: [] for t in WO.TSIS}
    for c in not_reported:
        t = codes[c].get("tsi") or "—"
        by_tsi.setdefault(t, []).append(c)
    for t in by_tsi:
        by_tsi[t].sort(key=lambda c: (codes[c].get("isd") or c).lower())

    day_s  = "%02d %s %d" % (day.day, day.strftime("%b"), day.year)
    hdr = ["🏁 <b>Competition Sales — NOT Reported · FTD</b>",
           "<i>%s · as of %s IST</i>" % (day_s, now_ist.strftime("%I:%M %p").lstrip("0")),
           "━" * 20]

    total = len(not_reported)
    if total == 0:
        hdr.append("✅ All eligible ISDs have reported competition sales for today.")
        return "\n".join(hdr)

    tallies = " · ".join("%s %d" % (t, len(by_tsi.get(t, []))) for t in WO.TSIS)
    hdr.append("Not reported: <b>%d</b> ISDs" % total)
    hdr.append(tallies)
    L = hdr + [""]
    order = WO.TSIS + [t for t in by_tsi if t not in WO.TSIS]
    for t in order:
        lst = by_tsi.get(t) or []
        if not lst:
            continue
        L.append("<b>%s</b> — %d" % (esc(t), len(lst)))
        for c in lst:
            info = codes[c]
            line = "   • %s" % esc(info.get("isd") or c)
            if info.get("store"):
                line += " — %s" % esc(info["store"])
            L.append(line)
        L.append("")
    if match_mode == "name":
        L.append("<i>Matched by ISD name (report has no code) — verify any unexpected name.</i>")
    return "\n".join(L).rstrip()

# ── main ─────────────────────────────────────────────────────────────────────

def main():
    _ovr = env("TODAY_OVERRIDE")
    now_ist = datetime.now(IST)
    day = datetime.strptime(_ovr, "%Y-%m-%d").date() if _ovr else now_ist.date()
    print("[COMP] FTD date:", day)

    codes = WO.load_codes()
    print("[COMP] roster ISDs:", len(codes))

    comp_xlsx = env("COMP_XLSX")
    path = Path(comp_xlsx) if comp_xlsx else comp_download(day, "_comp_dl")
    rows = _read_comp_xlsx(path)
    print("[COMP] competition rows parsed:", len(rows))

    reported, mode = reporters_for(day, rows, codes)
    print("[COMP] reported today:", len(reported), "| match mode:", mode)

    excl = excluded_today(day)

    not_reported = [c for c, info in codes.items()
                    if not AR._is_inactive(info)
                    and c not in excl
                    and c not in reported]
    print("[COMP] NOT reported:", len(not_reported))

    msg = build_message(not_reported, codes, day, now_ist, mode)
    WO.send_telegram(msg)
    print("[COMP] done.")


if __name__ == "__main__":
    main()
